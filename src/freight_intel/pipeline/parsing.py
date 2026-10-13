"""Step 1: turn a file (PDF, Excel, email) into plain text for the LLM.

No AI here. The aim is to keep what a human would see, in reading order, and to
say so honestly when we can't (for example a scanned PDF with no text layer).
"""
import email
import re
from dataclasses import dataclass, field
from email import policy
from html.parser import HTMLParser
from pathlib import Path


@dataclass
class ParsedDocument:
    source_format: str  # "pdf" | "xlsx" | "eml"
    text: str
    page_count: int = 1
    warnings: list[str] = field(default_factory=list)


def parse_document(path: str | Path) -> ParsedDocument:
    path = Path(path)
    suffix = path.suffix.lower().lstrip(".")
    if suffix == "pdf":
        return _parse_pdf(path)
    if suffix in ("xlsx", "xlsm"):
        return _parse_xlsx(path)
    if suffix == "eml":
        return _parse_eml(path)
    raise ValueError(f"unsupported file type: {path.suffix!r}")


def _parse_pdf(path: Path) -> ParsedDocument:
    import pdfplumber  # imported here so the other formats work without it

    pages = []
    with pdfplumber.open(path) as pdf:
        for page in pdf.pages:
            pages.append(page.extract_text() or "")
    text = "\n\n".join(pages).strip()
    warnings = []
    if not text:
        # A scanned PDF is just a picture. OCR is a later, optional step.
        warnings.append("no text layer found (scanned PDF?); OCR not implemented")
    return ParsedDocument("pdf", text, page_count=max(len(pages), 1), warnings=warnings)


def _cell_text(cell) -> str:
    """Show a cell the way Excel would, not the raw stored value."""
    v = cell.value
    if isinstance(v, bool):
        return str(v)
    if isinstance(v, (int, float)):
        if "0.00" in (cell.number_format or ""):
            return f"{v:,.2f}"  # number format '#,##0.00' => displays 2,960.00
        return str(int(v)) if float(v).is_integer() else repr(v)
    if hasattr(v, "date"):
        return v.date().isoformat()
    return str(v).strip()


def _parse_xlsx(path: Path) -> ParsedDocument:
    from openpyxl import load_workbook

    wb = load_workbook(path, data_only=True)
    lines: list[str] = []
    for ws in wb.worksheets:
        lines.append(f"[Sheet: {ws.title}]")
        for row in ws.iter_rows():
            # Empty cells (including the covered part of a merged cell) vanish.
            cells = [_cell_text(c) for c in row if c.value is not None and str(c.value).strip()]
            if cells:
                lines.append(" | ".join(cells))
    return ParsedDocument("xlsx", "\n".join(lines).strip(), page_count=len(wb.worksheets))


class _TextOnly(HTMLParser):
    def __init__(self):
        super().__init__()
        self.parts: list[str] = []

    def handle_data(self, data):
        self.parts.append(data)


def _html_to_text(html: str) -> str:
    parser = _TextOnly()
    parser.feed(html)
    return re.sub(r"\n{3,}", "\n\n", "".join(parser.parts)).strip()


def _parse_eml(path: Path) -> ParsedDocument:
    msg = email.message_from_bytes(path.read_bytes(), policy=policy.default)
    warnings = []
    body_part = msg.get_body(preferencelist=("plain", "html"))
    body = body_part.get_content() if body_part else ""
    if body_part is not None and body_part.get_content_type() == "text/html":
        body = _html_to_text(body)
    if any(True for _ in msg.iter_attachments()):
        warnings.append("email has attachments; attachments are not parsed yet")
    header = "\n".join(f"{k}: {msg[k]}" for k in ("From", "To", "Subject", "Date") if msg[k])
    return ParsedDocument("eml", f"{header}\n\n{body}".strip(), warnings=warnings)
