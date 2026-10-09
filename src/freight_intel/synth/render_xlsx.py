"""Excel renderer (openpyxl). Deliberately a bit untidy, like real spreadsheets:
a merged title cell, details in a loose key/value block, and amounts that are
real numbers for some carriers but text such as '1.234,56' for others."""
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from freight_intel.synth import domain
from freight_intel.synth.records import Case
from freight_intel.synth.view import DocumentView

NUMERIC_AMOUNTS = {"us"}  # carriers whose number style is 'us' store numbers; others store text


def render_xlsx(view: DocumentView, case: Case, path: Path) -> None:
    carrier = view.carrier
    wb = Workbook()
    ws = wb.active
    ws.title = view.title[:30].replace("/", "-")
    width = len(view.columns)
    bold = Font(bold=True)

    ws["A1"] = f"{carrier.name} - {view.title}"
    ws["A1"].font = Font(bold=True, size=14)
    if carrier.layout == "detailed":
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=width)  # merged cells = extraction pain

    row = 3
    for key, value in view.meta:
        ws.cell(row=row, column=1, value=key).font = bold
        ws.cell(row=row, column=2, value=value)
        row += 1
    row += 1  # blank spacer row

    for col, name in enumerate(view.columns, start=1):
        cell = ws.cell(row=row, column=col, value=name)
        cell.font = bold
        cell.fill = PatternFill("solid", fgColor="DDDDDD")
    row += 1

    numeric = carrier.number_style in NUMERIC_AMOUNTS
    charges = case.quote.charges if view.kind == "quote" else case.invoice.charges
    total_minor = case.quote.total_minor if view.kind == "quote" else case.invoice.stated_total_minor
    amounts = [c.amount_minor for c in charges] + [total_minor]

    for text_row, minor in zip(view.rows + [view.total], amounts):
        for col, text in enumerate(text_row, start=1):
            if col - 1 == view.amount_col and numeric:
                # A real number with a display format. The currency is NOT in the
                # cell; the extractor must read it from the 'Currency' line above.
                cell = ws.cell(row=row, column=col, value=minor / 100)
                cell.number_format = "#,##0.00"
            else:
                cell = ws.cell(row=row, column=col, value=text)
            if text_row is view.total:
                cell.font = bold
        row += 1

    row += 1
    for line in view.remarks:
        ws.cell(row=row, column=1, value=line).alignment = Alignment(wrap_text=False)
        row += 1

    ws.column_dimensions["A"].width = 36
    for letter in "BCD"[: max(width - 1, 1)]:
        ws.column_dimensions[letter].width = 22
    wb.save(path)
