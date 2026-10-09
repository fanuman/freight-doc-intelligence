"""Tests for the synthetic data generator.

These need no database. They check that the ground truth is internally
consistent (a dataset whose answers are wrong would make every later
evaluation meaningless) and that the renderers actually print that truth.
"""
import email
from collections import Counter
from dataclasses import asdict

import pytest

from freight_intel.synth import domain
from freight_intel.synth import scenario as sc
from freight_intel.synth.scenario import build_case, build_dataset

SEED = 42
N = 45  # the default dataset: 5 of each plan slot


@pytest.fixture(scope="module")
def dataset():
    return build_dataset(N, SEED)


# --- money formatting ---------------------------------------------------------

@pytest.mark.parametrize("minor, style, expected", [
    (123456, "us", "1,234.56"),
    (123456, "eu", "1.234,56"),
    (123456, "space", "1 234.56"),
    (5, "us", "0.05"),
    (100, "eu", "1,00"),
    (123456789, "eu", "1.234.567,89"),
    (-2500, "us", "-25.00"),
    (0, "us", "0.00"),
])
def test_format_amount(minor, style, expected):
    assert domain.format_amount(minor, style) == expected


def test_format_money_styles():
    assert domain.format_money(123456, "USD", "us", "prefix") == "USD 1,234.56"
    assert domain.format_money(123456, "EUR", "eu", "suffix") == "1.234,56 EUR"
    assert domain.format_money(123456, "EUR", "us", "symbol") == "€1,234.56"


def test_currency_conversion_is_integer_and_rounds_half_up():
    assert domain.convert_minor(100, "USD", "EUR") == 92
    assert domain.convert_minor(5, "USD", "EUR") == 5  # 4.6 -> 5
    assert isinstance(domain.convert_minor(123457, "EUR", "USD"), int)


# --- determinism and coverage ---------------------------------------------------

def test_same_seed_gives_identical_cases():
    assert asdict(build_case(7, SEED)) == asdict(build_case(7, SEED))


def test_different_seed_gives_different_cases():
    assert asdict(build_case(7, 1)) != asdict(build_case(7, 2))


def test_every_error_kind_appears_equally_often(dataset):
    counts = Counter(c.error_kind for c in dataset)
    assert set(counts) == set(sc.ERROR_PLAN)
    assert all(counts[k] == 5 for k in set(sc.ERROR_PLAN) - {sc.CLEAN})
    assert counts[sc.CLEAN] == 10


def test_all_four_carriers_are_used(dataset):
    assert {c.quote.carrier_id for c in dataset} == {c.carrier_id for c in domain.CARRIERS}


def test_some_invoices_carry_a_prompt_injection(dataset):
    injected = [c for c in dataset if c.prompt_injection]
    assert injected and all(c.injection_text for c in injected)
    assert not any(c.injection_text for c in dataset if not c.prompt_injection)


# --- ground truth is internally consistent ---------------------------------------

def test_all_amounts_are_integers(dataset):
    for case in dataset:
        for c in case.quote.charges + case.invoice.charges:
            assert type(c.amount_minor) is int and c.amount_minor > 0
        assert type(case.invoice.stated_total_minor) is int


def test_printed_total_matches_lines_except_arithmetic_error(dataset):
    for case in dataset:
        same = case.invoice.stated_total_minor == case.invoice.computed_total_minor
        assert same == (case.error_kind != sc.ARITHMETIC_ERROR), case.case_id


def test_only_expired_cases_are_invoiced_after_validity(dataset):
    for case in dataset:
        expired = case.invoice.invoice_date > case.quote.valid_until
        assert expired == (case.error_kind == sc.EXPIRED_QUOTE), case.case_id


def test_only_currency_mismatch_cases_change_currency(dataset):
    for case in dataset:
        changed = case.invoice.currency != case.quote.currency
        assert changed == (case.error_kind == sc.CURRENCY_MISMATCH), case.case_id


def test_clean_and_rounding_cases_expect_no_discrepancies(dataset):
    for case in dataset:
        if case.error_kind in (sc.CLEAN, sc.ROUNDING_NOISE):
            assert case.expected == [], case.case_id


def test_every_seeded_error_has_exactly_the_expected_discrepancy(dataset):
    for case in dataset:
        if case.error_kind in (sc.CLEAN, sc.ROUNDING_NOISE):
            continue
        assert [d.type for d in case.expected] == [case.error_kind], case.case_id


def test_overcharge_variance_matches_the_two_documents(dataset):
    for case in (c for c in dataset if c.error_kind == sc.OVERCHARGE):
        d = case.expected[0]
        quoted = next(c for c in case.quote.charges if c.code == d.charge_code)
        invoiced = next(c for c in case.invoice.charges if c.code == d.charge_code)
        assert d.quoted_minor == quoted.amount_minor
        assert d.invoiced_minor == invoiced.amount_minor
        assert d.variance_minor == invoiced.amount_minor - quoted.amount_minor > 0
        # well above the 0.5% tolerance, so a correct engine must flag it
        assert d.variance_minor * 200 > quoted.amount_minor


def test_unquoted_charge_is_really_absent_from_the_quote(dataset):
    for case in (c for c in dataset if c.error_kind == sc.UNQUOTED_CHARGE):
        d = case.expected[0]
        assert d.quoted_minor is None
        assert (d.charge_code, d.invoiced_minor) in [
            (c.code, c.amount_minor) for c in case.invoice.charges]
        quoted_labels = {c.raw_label for c in case.quote.charges}
        assert not any(c.raw_label in quoted_labels and c.code == d.charge_code
                       for c in case.invoice.charges if c.amount_minor == d.invoiced_minor)


def test_missing_charge_is_present_on_quote_only(dataset):
    for case in (c for c in dataset if c.error_kind == sc.MISSING_CHARGE):
        d = case.expected[0]
        assert d.charge_code in {c.code for c in case.quote.charges}
        assert d.charge_code not in {c.code for c in case.invoice.charges}
        assert d.variance_minor == -d.quoted_minor


def test_rounding_noise_is_one_minor_unit(dataset):
    for case in (c for c in dataset if c.error_kind == sc.ROUNDING_NOISE):
        diffs = [i.amount_minor - q.amount_minor
                 for q, i in zip(case.quote.charges, case.invoice.charges)]
        assert sorted(diffs) == [0] * (len(diffs) - 1) + [1]


def test_invoice_labels_differ_from_quote_labels(dataset):
    """Quote and invoice spell charges differently, so matching can't be string equality."""
    for case in dataset:
        q = {c.code: c.raw_label for c in case.quote.charges}
        differing = sum(1 for c in case.invoice.charges if c.code in q and c.raw_label != q[c.code])
        assert differing >= 3, case.case_id


# --- rendering -----------------------------------------------------------------------

@pytest.fixture(scope="module")
def rendered(tmp_path_factory, dataset):
    pytest.importorskip("reportlab")
    pytest.importorskip("openpyxl")
    from freight_intel.synth.writer import write_dataset

    out = tmp_path_factory.mktemp("synthetic")
    write_dataset(N, SEED, out)
    return out


def test_manifest_and_files_exist(rendered, dataset):
    import json

    manifest = json.loads((rendered / "manifest.json").read_text())
    assert manifest["count"] == N and len(manifest["cases"]) == N
    for entry in manifest["cases"]:
        for key in ("quote_file", "invoice_file", "truth_file"):
            assert (rendered / entry[key]).stat().st_size > 0, entry[key]


def test_pdfs_are_real_pdfs_containing_the_printed_amounts(rendered, dataset):
    checked = 0
    for case in dataset:
        if case.invoice_format != "pdf":
            continue
        raw = (rendered / "cases" / case.case_id / "invoice.pdf").read_bytes()
        assert raw.startswith(b"%PDF")
        carrier = next(c for c in domain.CARRIERS if c.carrier_id == case.quote.carrier_id)
        for charge in case.invoice.charges:
            assert domain.format_amount(charge.amount_minor, carrier.number_style).encode() in raw
        assert domain.format_amount(case.invoice.stated_total_minor, carrier.number_style).encode() in raw
        checked += 1
    assert checked > 10


def test_xlsx_contains_the_printed_total(rendered, dataset):
    from openpyxl import load_workbook

    checked = 0
    for case in dataset:
        if case.invoice_format != "xlsx":
            continue
        ws = load_workbook(rendered / "cases" / case.case_id / "invoice.xlsx").active
        cells = [c.value for row in ws.iter_rows() for c in row if c.value is not None]
        carrier = next(c for c in domain.CARRIERS if c.carrier_id == case.quote.carrier_id)
        if carrier.number_style == "us":
            assert case.invoice.stated_total_minor / 100 in cells
        else:
            assert domain.format_amount(case.invoice.stated_total_minor, carrier.number_style) in cells
        checked += 1
    assert checked > 5


def test_eml_quotes_parse_and_show_the_ocean_freight_amount(rendered, dataset):
    checked = 0
    for case in dataset:
        if case.quote_format != "eml":
            continue
        raw = (rendered / "cases" / case.case_id / "quote.eml").read_bytes()
        msg = email.message_from_bytes(raw)
        assert case.quote.quote_ref in msg["Subject"]
        body = msg.get_payload(decode=True).decode()
        carrier = next(c for c in domain.CARRIERS if c.carrier_id == case.quote.carrier_id)
        ocean = next(c for c in case.quote.charges if c.code == "OCEAN_FREIGHT")
        assert domain.format_amount(ocean.amount_minor, carrier.number_style) in body
        checked += 1
    assert checked > 5


def test_injection_text_is_in_the_rendered_invoice_but_not_in_truth_charges(rendered, dataset):
    for case in (c for c in dataset if c.prompt_injection and c.invoice_format == "pdf"):
        raw = (rendered / "cases" / case.case_id / "invoice.pdf").read_bytes()
        assert case.injection_text.split(".")[0].encode()[:30] in raw
